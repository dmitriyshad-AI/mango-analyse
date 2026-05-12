from __future__ import annotations

import argparse
import csv
import json
import re
from copy import copy
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from mango_mvp.insights.sanitizers import (
    flag_booleans,
    flags_to_text,
    has_any_safety_risk,
    sanitize_answer,
    sanitize_customer_text,
)


POSITIVE_OUTCOMES = {"won_paid_or_active", "payment_pending"}
OPPORTUNITY_OUTCOMES = {"reopen_or_follow_up_opportunity", "in_progress_or_undecided", "known_student_or_lead"}
RETENTION_OUTCOMES = {"service_or_existing_context", "existing_client_service_not_new_sale"}
NEGATIVE_OUTCOMES = {"lost_or_refused", "closed_lost_valid", "churn_or_refused_after_activity"}
SIGNAL_LABELS_RU = {
    "price_question": "Вопрос о цене",
    "price_objection": "Возражение по цене",
    "discount_or_installment_question": "Скидка / рассрочка",
    "schedule_question": "Расписание / время",
    "format_question_online_offline": "Формат онлайн/очно",
    "location_question": "Локация / филиал",
    "teacher_question": "Вопрос о преподавателе",
    "program_question": "Программа / курс",
    "level_fit_question": "Подходит ли уровень",
    "exam_or_olympiad_goal": "ЕГЭ/ОГЭ/олимпиады",
    "trust_question": "Доверие / качество",
    "competitor_comparison": "Сравнение с конкурентами",
    "child_motivation_concern": "Мотивация ребенка",
    "parent_decision_delay": "Родитель откладывает решение",
    "spouse_or_family_approval": "Нужно согласование в семье",
    "not_relevant_now": "Сейчас не актуально",
    "already_learning_elsewhere": "Уже учатся в другом месте",
    "ready_to_pay": "Готовность оплатить",
    "materials_request": "Просьба прислать материалы",
    "callback_request": "Просьба перезвонить",
    "technical_or_access_issue": "Технический вопрос / доступ",
    "existing_client_progress": "Прогресс действующего клиента",
    "complaint_or_service_risk": "Жалоба / сервисный риск",
    "payment_or_contract_service": "Оплата / договор / документы",
    "unknown": "Неясный сигнал",
}
STAGE_LABELS_RU = {
    "new_request": "Новая заявка",
    "discovery": "Выявление потребности",
    "offer_explained": "Объяснение предложения",
    "price_discussion": "Обсуждение цены",
    "objection_handling": "Работа с возражением",
    "materials_sent": "Материалы / КП",
    "decision_wait": "Клиент думает",
    "payment_intent": "Намерение оплатить",
    "paid_or_enrolled": "Оплатил / записан",
    "existing_client_service": "Сервис действующего клиента",
    "reactivation": "Реактивация",
    "lost_or_stalled": "Потерян / завис",
    "retention_or_expansion": "Удержание / допродажа",
    "unknown": "Неясная стадия",
}
ANSWER_PATTERN_LABELS_RU = {
    "vague_or_missing_next_step": "Нет точного следующего шага",
    "service_resolution_or_escalation": "Сервисный вопрос взят в работу",
    "handoff_or_transfer": "Перевод / передача другому сотруднику",
    "program_fit_explained": "Программа или формат объяснены",
    "program_fit_not_explained": "Программа или формат не объяснены",
    "next_step_or_materials_promised": "Обещаны материалы или следующий контакт",
    "no_live_contact_or_voicemail": "Не было живого контакта / автоответчик",
    "price_payment_handled_with_value_or_instruction": "Цена/оплата объяснены через ценность или инструкцию",
    "price_payment_answer_too_weak": "Слабый ответ на цену/оплату",
    "diagnostic_questions": "Менеджер задавал диагностические вопросы",
    "decision_delay_acknowledged": "Отложенное решение принято без дожима",
    "service_answer_without_closure": "Сервисный ответ без закрытия вопроса",
    "logistics_answer_without_commitment": "Логистика без четкого обязательства",
    "generic_answer": "Общий ответ без выраженного паттерна",
}
OUTCOME_GROUP_LABELS_RU = {
    "paid_or_payment_path": "Оплата или путь к оплате",
    "follow_up_opportunity": "Есть возможность повторного контакта / реактивации",
    "retention_or_service": "Удержание / сервис",
    "lost_or_churn": "Потеря / отток",
    "manual_or_mixed": "Смешанный случай / ручная проверка",
    "unknown_or_other": "Неясный итог",
}
FINAL_OUTCOME_LABELS_RU = {
    "won_paid_or_active": "Оплатил / действующий клиент",
    "payment_pending": "Есть путь к оплате",
    "reopen_or_follow_up_opportunity": "Нужен повторный контакт / реактивация",
    "in_progress_or_undecided": "В процессе / клиент думает",
    "known_student_or_lead": "Известный ученик или лид",
    "service_or_existing_context": "Сервисный контакт действующего клиента",
    "existing_client_service_not_new_sale": "Действующий клиент, не новая продажа",
    "lost_or_refused": "Отказ / потеря",
    "closed_lost_valid": "Закрыто как проигрыш обоснованно",
    "churn_or_refused_after_activity": "Отток / отказ после активности",
    "mixed_outcome_manual_review": "Смешанный итог, нужна ручная проверка",
    "unknown": "Неясный итог",
}
BOT_STATUS_LABELS_RU = {
    "ready_for_bot_draft": "Можно брать как черновик для бота",
    "needs_rop_validation": "Нужна проверка РОПом",
    "not_ready": "Пока не готово",
    "not_ready_needs_llm_refresh": "Пока не готово: нужен live LLM-refresh",
    "not_ready_sanitizer_blocked": "Пока не готово: sanitizer нашел риск",
    "exclude_no_dialogue": "Исключить: не было диалога",
    "manager_process_only": "Только процесс менеджера, не база бота",
}
USEFULNESS_LABELS_RU = {
    "playbook_candidate": "Кандидат в базу лучших ответов",
    "revenue_leakage_risk": "Риск потери выручки",
    "service_retention_learning": "Урок для удержания/сервиса",
    "process_fix_needed": "Нужна правка процесса",
    "coaching_needed": "Нужен разбор с менеджером",
    "needs_llm_refresh": "Нужно живое LLM-ревью",
    "useful_context": "Полезный контекст",
}
REVIEW_TRUST_LABELS_RU = {
    "trusted_llm_review": "Доверенное live/GPT-ревью",
    "needs_live_llm_refresh": "Нужен live LLM-refresh перед использованием",
}
BOT_SAFETY_LABELS_RU = {
    "safe_no_changes": "Безопасно без замен",
    "safe_with_placeholders": "Безопасно после sanitization",
    "blocked_unresolved_safety_risk": "Заблокировано: остался safety-риск",
    "fixpoint_not_reached": "Заблокировано: sanitizer не достиг стабильного результата",
    "empty": "Пустой ответ",
}
BOT_BLOCKING_SAFETY_STATUSES = {"blocked_unresolved_safety_risk", "fixpoint_not_reached"}


@dataclass(frozen=True)
class KnowledgeBaseConfig:
    project_root: Path
    reviews_csv: Path
    out_root: Path
    min_group_count: int = 2
    top_examples: int = 120


def build_sales_insight_knowledge_base(config: KnowledgeBaseConfig) -> dict[str, Any]:
    out_root = config.out_root.resolve()
    out_root.mkdir(parents=True, exist_ok=True)
    reviews = [enrich_review_row(row) for row in read_csv(config.reviews_csv)]
    trusted_reviews = [row for row in reviews if is_trusted_llm_review(row)]

    summary = build_summary(config, reviews)
    signal_summary = build_signal_summary(trusted_reviews)
    pattern_matrix = build_pattern_matrix(trusted_reviews, config.min_group_count)
    best_answers = build_best_answer_candidates(trusted_reviews, config.top_examples)
    rop_coaching = build_rop_coaching_queue(trusted_reviews, config.top_examples)
    bot_seeds = build_bot_knowledge_seeds(trusted_reviews, config.top_examples)
    manager_summary = build_manager_summary(trusted_reviews)
    outcome_lens = build_outcome_lens(trusted_reviews)
    llm_refresh_queue = build_llm_refresh_queue(reviews)
    rop_brief = build_rop_brief(summary, signal_summary, pattern_matrix, rop_coaching, bot_seeds, manager_summary)

    outputs = write_outputs(
        out_root,
        summary,
        rop_brief,
        signal_summary,
        pattern_matrix,
        best_answers,
        rop_coaching,
        bot_seeds,
        manager_summary,
        outcome_lens,
        llm_refresh_queue,
        reviews,
    )
    summary["outputs"] = {key: str(path) for key, path in outputs.items()}
    (out_root / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def enrich_review_row(row: dict[str, Any]) -> dict[str, Any]:
    enriched = dict(row)
    score = clamp_int(row.get("overall_quality_score"), 0, 100, 0)
    signal = clean(row.get("llm_customer_signal_type")) or "unknown"
    stage = clean(row.get("llm_hidden_sales_stage")) or "unknown"
    outcome = clean(row.get("final_outcome_label")) or "unknown"
    trusted = is_trusted_llm_review(row)
    manager_answer = sanitize_answer(row.get("ideal_answer_example"), mode="manager")
    bot_answer = sanitize_answer(row.get("ideal_answer_example"), mode="bot")
    customer_question = sanitize_customer_text(row.get("customer_question"))
    customer_quote = sanitize_customer_text(row.get("customer_quote"))
    manager_quote = sanitize_customer_text(row.get("manager_quote"))
    sanitizer_flags = tuple(
        dict.fromkeys((*manager_answer.flags, *bot_answer.flags, *customer_question.flags, *customer_quote.flags, *manager_quote.flags))
    )
    enriched["overall_quality_score"] = score
    enriched["quality_band"] = quality_band(score)
    enriched["outcome_group"] = outcome_group(outcome)
    enriched["answer_pattern"] = classify_answer_pattern(row)
    enriched["review_trust_status"] = "trusted_llm_review" if trusted else "needs_live_llm_refresh"
    enriched["ideal_answer_manager_sanitized"] = manager_answer.text
    enriched["bot_safe_answer"] = bot_answer.text
    enriched["customer_question_sanitized"] = customer_question.text
    enriched["customer_quote_sanitized"] = customer_quote.text
    enriched["manager_quote_sanitized"] = manager_quote.text
    enriched["sanitizer_flags"] = flags_to_text(sanitizer_flags)
    enriched.update(flag_booleans(sanitizer_flags))
    enriched["bot_safety_status"] = bot_answer.status
    enriched["bot_sanitizer_pass_count"] = bot_answer.pass_count
    enriched["bot_sanitizer_fixpoint_reached"] = "Да" if bot_answer.fixpoint_reached else "Нет"
    enriched["bot_safety_blocked"] = "Да" if has_any_safety_risk(bot_answer.text) else "Нет"
    if trusted:
        enriched["commercial_usefulness"] = commercial_usefulness(row, score, outcome)
        enriched["rop_action"] = rop_action_for_row(row, score)
        enriched["bot_seed_status"] = bot_seed_status(
            row,
            score,
            bot_answer=bot_answer.text,
            bot_safety_status=bot_answer.status,
            sanitizer_flags=sanitizer_flags,
        )
    else:
        enriched["commercial_usefulness"] = "needs_llm_refresh"
        enriched["rop_action"] = "Сначала запустить live LLM-refresh; до этого не использовать строку как пример для РОПа или бота."
        enriched["bot_seed_status"] = "not_ready_needs_llm_refresh"
    enriched["signal_ru"] = signal_label_ru(signal)
    enriched["stage_ru"] = stage_label_ru(stage)
    enriched["outcome_group_ru"] = outcome_group_label_ru(enriched["outcome_group"])
    enriched["final_outcome_ru"] = final_outcome_label_ru(outcome)
    enriched["answer_pattern_ru"] = answer_pattern_label_ru(enriched["answer_pattern"])
    enriched["commercial_usefulness_ru"] = usefulness_label_ru(enriched["commercial_usefulness"])
    enriched["bot_seed_status_ru"] = bot_status_label_ru(enriched["bot_seed_status"])
    enriched["review_trust_status_ru"] = review_trust_label_ru(enriched["review_trust_status"])
    enriched["bot_safety_status_ru"] = bot_safety_label_ru(enriched["bot_safety_status"])
    enriched["data_scope_note"] = "Оценка только по звонкам: мессенджеры и почта в этом слое не учтены."
    enriched["signal_stage_key"] = f"{signal}::{stage}"
    enriched["signal_pattern_key"] = f"{signal}::{enriched['answer_pattern']}"
    return enriched


def is_trusted_llm_review(row: dict[str, Any]) -> bool:
    provider = clean(row.get("provider"))
    review_source = clean(row.get("review_source"))
    if provider == "dry_run":
        return False
    if "deterministic_fallback" in review_source or "needs_llm_refresh" in review_source:
        return False
    return bool(provider)


def quality_band(score: int) -> str:
    if score >= 75:
        return "high"
    if score >= 55:
        return "medium"
    return "low"


def outcome_group(outcome: Any) -> str:
    label = clean(outcome)
    if label in POSITIVE_OUTCOMES:
        return "paid_or_payment_path"
    if label in OPPORTUNITY_OUTCOMES:
        return "follow_up_opportunity"
    if label in RETENTION_OUTCOMES:
        return "retention_or_service"
    if label in NEGATIVE_OUTCOMES:
        return "lost_or_churn"
    if "manual" in label or "mixed" in label:
        return "manual_or_mixed"
    return "unknown_or_other"


def classify_answer_pattern(row: dict[str, Any]) -> str:
    signal = clean(row.get("llm_customer_signal_type"))
    stage = clean(row.get("llm_hidden_sales_stage"))
    answer = clean(row.get("manager_answer"))
    risk = clean(row.get("risk_flags"))
    missed = clean(row.get("what_manager_missed"))
    text = " ".join([answer, risk, missed, clean(row.get("what_manager_did_well"))]).lower()

    if re.search(
        r"абонент|недозвон|автоответчик|нет живого|контакт не состоя|не было диалога|voicemail|no_live|"
        r"продолжение следует|субтитры сделал|редактор субтитров|thank you for watching|norske lagerforskning",
        text,
    ):
        return "no_live_contact_or_voicemail"
    if re.search(r"перев[её]л|соедин|передал[аи]?|коллег|администратор|другой менеджер", text):
        return "handoff_or_transfer"
    if re.search(r"нет точн|без точн|не зафиксирован|размыт|может остыть|сам перезвон|без даты|без срока", text):
        return "vague_or_missing_next_step"
    if signal in {"complaint_or_service_risk", "technical_or_access_issue"} or stage == "existing_client_service":
        if re.search(r"провер|уточн|передам|продублир|отправ|эскал|исправ|поддерж", text):
            return "service_resolution_or_escalation"
        return "service_answer_without_closure"
    if signal in {"price_question", "price_objection", "discount_or_installment_question", "payment_or_contract_service"}:
        if re.search(r"ценност|результат|выгод|программ|преподав|уровн|рассроч|скидк|маткапитал|ссылк|сч[её]т|квитанц|оплат", text):
            return "price_payment_handled_with_value_or_instruction"
        return "price_payment_answer_too_weak"
    if signal in {"program_question", "format_question_online_offline", "location_question", "teacher_question", "level_fit_question", "exam_or_olympiad_goal"}:
        if re.search(r"групп|курс|программ|формат|очно|онлайн|уровн|тест|преподав|егэ|огэ|олимпиад|филиал|локац", text):
            return "program_fit_explained"
        return "program_fit_not_explained"
    if signal in {"schedule_question", "callback_request", "materials_request"}:
        if re.search(r"отправ|пришл|перезвон|верн[её]мся|соглас|срок|дат|врем|распис", text):
            return "next_step_or_materials_promised"
        return "logistics_answer_without_commitment"
    if re.search(r"подума|решени|обсуд|совет|жена|семь", text):
        return "decision_delay_acknowledged"
    if re.search(r"вопрос|уточн|подскаж|диагност", text):
        return "diagnostic_questions"
    return "generic_answer"


def commercial_usefulness(row: dict[str, Any], score: int, outcome: Any) -> str:
    group = outcome_group(outcome)
    pattern = classify_answer_pattern(row)
    if pattern == "no_live_contact_or_voicemail":
        return "process_fix_needed"
    if score >= 75 and group in {"paid_or_payment_path", "follow_up_opportunity"}:
        return "playbook_candidate"
    if score < 55 and group in {"lost_or_churn", "follow_up_opportunity"}:
        return "revenue_leakage_risk"
    if score < 55:
        return "coaching_needed"
    if group == "retention_or_service":
        return "service_retention_learning"
    if pattern in {"no_live_contact_or_voicemail", "vague_or_missing_next_step"}:
        return "process_fix_needed"
    return "useful_context"


def rop_action_for_row(row: dict[str, Any], score: int) -> str:
    signal = clean(row.get("llm_customer_signal_type"))
    pattern = classify_answer_pattern(row)
    if pattern == "no_live_contact_or_voicemail":
        return "Ввести правило: после недозвона фиксировать точное время повторного звонка и альтернативный канал."
    if pattern == "vague_or_missing_next_step":
        return "Тренировать обязательный конкретный следующий шаг: дата, время, канал, ответственный."
    if pattern in {"price_payment_answer_too_weak", "price_payment_handled_with_value_or_instruction"} and score < 75:
        return "Разобрать блок цены: сначала ценность/результат, затем условия оплаты, затем следующий шаг."
    if signal in {"complaint_or_service_risk", "technical_or_access_issue"}:
        return "Поставить восстановление сервиса: признать проблему, назвать срок решения, проконтролировать закрытие."
    if pattern in {"program_fit_not_explained", "program_fit_explained"} and score < 75:
        return "Усилить квалификацию: цель, уровень, формат, ограничения, затем точная рекомендация курса."
    if score >= 75:
        return "Использовать как пример хорошей реакции в базе лучших ответов после ручной проверки РОПом."
    return "Проверить с менеджером: что было потребностью клиента и какой следующий шаг должен быть в CRM."


def bot_seed_status(
    row: dict[str, Any],
    score: int,
    *,
    bot_answer: str | None = None,
    bot_safety_status: str | None = None,
    sanitizer_flags: tuple[str, ...] | list[str] = (),
) -> str:
    signal = clean(row.get("llm_customer_signal_type"))
    ideal = clean(bot_answer if bot_answer is not None else row.get("ideal_answer_example"))
    confidence = clamp_float(row.get("extraction_confidence"), 0.0, 1.0, 0.0)
    pattern = classify_answer_pattern(row)
    if pattern == "no_live_contact_or_voicemail":
        return "exclude_no_dialogue"
    if bot_safety_status in BOT_BLOCKING_SAFETY_STATUSES:
        return "not_ready_sanitizer_blocked"
    if sanitizer_flags:
        return "needs_rop_validation" if score >= 60 and len(ideal) >= 60 else "not_ready"
    if signal in {"callback_request", "materials_request"} and score < 65:
        return "manager_process_only"
    if score >= 75 and confidence >= 0.65 and len(ideal) >= 80:
        return "ready_for_bot_draft"
    if score >= 60 and len(ideal) >= 60:
        return "needs_rop_validation"
    return "not_ready"


def build_summary(config: KnowledgeBaseConfig, reviews: list[dict[str, Any]]) -> dict[str, Any]:
    scores = [clamp_int(row.get("overall_quality_score"), 0, 100, 0) for row in reviews]
    trusted_count = sum(1 for row in reviews if is_trusted_llm_review(row))
    needs_refresh_count = len(reviews) - trusted_count
    bot_ready_rows = [row for row in reviews if row.get("bot_seed_status") in {"ready_for_bot_draft", "needs_rop_validation"}]
    bot_safe_rows = [row for row in bot_ready_rows if clean(row.get("bot_safe_answer")) and row.get("bot_safety_status") not in BOT_BLOCKING_SAFETY_STATUSES]
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_reviews_csv": str(config.reviews_csv.resolve()),
        "config": {
            "min_group_count": config.min_group_count,
            "top_examples": config.top_examples,
        },
        "totals": {
            "reviews": len(reviews),
            "unique_phones": len({clean(row.get("phone")) for row in reviews if clean(row.get("phone"))}),
            "unique_managers": len({clean(row.get("manager_name")) for row in reviews if clean(row.get("manager_name"))}),
            "signals": len({clean(row.get("llm_customer_signal_type")) for row in reviews if clean(row.get("llm_customer_signal_type"))}),
            "answer_patterns": len({clean(row.get("answer_pattern")) for row in reviews if clean(row.get("answer_pattern"))}),
        },
        "quality": {
            "avg_quality_score": round(sum(scores) / len(scores), 2) if scores else 0,
            "high_quality_count": sum(1 for score in scores if score >= 75),
            "low_quality_count": sum(1 for score in scores if score < 55),
        },
        "llm_review": {
            "trusted_llm_reviews": trusted_count,
            "needs_live_llm_refresh": needs_refresh_count,
            "trusted_share": round(trusted_count / len(reviews), 4) if reviews else 0,
            "by_provider": dict(Counter(clean(row.get("provider")) for row in reviews).most_common()),
            "by_review_source": dict(Counter(clean(row.get("review_source")) for row in reviews).most_common()),
            "by_review_trust_status": dict(Counter(clean(row.get("review_trust_status")) for row in reviews).most_common()),
        },
        "sanitizer": {
            "bot_ready_or_validation_rows": len(bot_ready_rows),
            "bot_safe_answer_rows": len(bot_safe_rows),
            "bot_safety_blocked": sum(1 for row in reviews if row.get("bot_safety_status") in BOT_BLOCKING_SAFETY_STATUSES),
            "by_bot_safety_status": dict(Counter(clean(row.get("bot_safety_status")) for row in reviews).most_common()),
            "by_bot_sanitizer_pass_count": dict(Counter(str(row.get("bot_sanitizer_pass_count") or "") for row in reviews).most_common()),
            "bot_fixpoint_not_reached": sum(1 for row in reviews if row.get("bot_sanitizer_fixpoint_reached") == "Нет"),
            "by_sanitizer_flag": dict(sanitizer_flag_counter(reviews).most_common()),
        },
        "counts": {
            "by_signal": dict(Counter(clean(row.get("llm_customer_signal_type")) for row in reviews).most_common()),
            "by_answer_pattern": dict(Counter(clean(row.get("answer_pattern")) for row in reviews).most_common()),
            "by_outcome_group": dict(Counter(clean(row.get("outcome_group")) for row in reviews).most_common()),
            "by_bot_seed_status": dict(Counter(clean(row.get("bot_seed_status")) for row in reviews).most_common()),
            "by_commercial_usefulness": dict(Counter(clean(row.get("commercial_usefulness")) for row in reviews).most_common()),
            "by_bot_safety_status": dict(Counter(clean(row.get("bot_safety_status")) for row in reviews).most_common()),
        },
        "audit_notes": [
            "Это агрегаты по live LLM-review выбранных sales moments; использовать как управленческую аналитику и материал для ручной проверки РОПом.",
            "Лучшие ответы, черновики бота, паттерны и очередь РОПа строятся только из доверенных live/GPT-ревью; строки dry_run/fallback вынесены в очередь LLM-refresh.",
            "Bot-safe ответы проходят deterministic sanitizer: брендовые ASR-артефакты, цены, скидки, сроки, возвраты, рассрочки и персональные данные заменяются безопасными формулировками.",
            "Важно: здесь учтены только звонки. Мессенджеры и почта не включены, поэтому слабый следующий шаг в звонке может быть закрыт менеджером в другом канале.",
            "Итог сделки показывает корреляцию/контекст, не причинность: нужен контроль источника, менеджера, курса и сезонности.",
            "Черновики для бота не являются готовой базой знаний: перед применением в Telegram-боте нужна ручная проверка РОПа/методиста.",
        ],
    }


def build_signal_summary(reviews: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for signal, group_rows in sorted(group_by(reviews, "llm_customer_signal_type").items()):
        scores = [clamp_int(row.get("overall_quality_score"), 0, 100, 0) for row in group_rows]
        rows.append(
            {
                "Сигнал клиента": signal_label_ru(signal),
                "Код сигнала": signal or "unknown",
                "Количество": len(group_rows),
                "Среднее качество": avg(scores),
                "Высокое качество": sum(1 for score in scores if score >= 75),
                "Низкое качество": sum(1 for score in scores if score < 55),
                "Доля оплат / пути к оплате": share(group_rows, "outcome_group", "paid_or_payment_path"),
                "Доля потерь / оттока": share(group_rows, "outcome_group", "lost_or_churn"),
                "Топ паттерны ответа": top_labeled_values((row.get("answer_pattern") for row in group_rows), answer_pattern_label_ru, 4),
                "Топ стадий": top_labeled_values((row.get("llm_hidden_sales_stage") for row in group_rows), stage_label_ru, 4),
                "Рекомендация РОПу": signal_level_recommendation(signal, group_rows),
                "Ограничение данных": "Только звонки; переписка в мессенджерах/почте не учтена.",
            }
        )
    return sorted(rows, key=lambda row: (-int(row["Количество"]), str(row["Сигнал клиента"])))


def build_pattern_matrix(reviews: list[dict[str, Any]], min_group_count: int) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in reviews:
        grouped[(clean(row.get("llm_customer_signal_type")), clean(row.get("answer_pattern")), clean(row.get("outcome_group")))].append(row)
    rows: list[dict[str, Any]] = []
    for (signal, pattern, outcome), group_rows in grouped.items():
        if len(group_rows) < max(1, min_group_count):
            continue
        scores = [clamp_int(row.get("overall_quality_score"), 0, 100, 0) for row in group_rows]
        example = best_row(group_rows)
        rows.append(
            {
                "Сигнал клиента": signal_label_ru(signal),
                "Код сигнала": signal,
                "Паттерн ответа": answer_pattern_label_ru(pattern),
                "Код паттерна": pattern,
                "Группа итога": outcome_group_label_ru(outcome),
                "Код группы итога": outcome,
                "Количество": len(group_rows),
                "Среднее качество": avg(scores),
                "Высокое качество": sum(1 for score in scores if score >= 75),
                "Низкое качество": sum(1 for score in scores if score < 55),
                "Лучший пример: оценка": example.get("overall_quality_score", ""),
                "Лучший пример вопрос": example.get("customer_question_sanitized", example.get("customer_question", "")),
                "Лучший идеальный ответ": example.get("ideal_answer_manager_sanitized", example.get("ideal_answer_example", "")),
                "Лучший безопасный ответ для бота": example.get("bot_safe_answer", ""),
                "Что делать": pattern_level_recommendation(pattern, outcome, avg(scores)),
            }
        )
    return sorted(rows, key=lambda row: (-int(row["Количество"]), str(row["Сигнал клиента"]), str(row["Паттерн ответа"])))


def build_best_answer_candidates(reviews: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    candidates = [
        row
        for row in reviews
        if is_trusted_llm_review(row)
        and clamp_int(row.get("overall_quality_score"), 0, 100, 0) >= 75
        and row.get("answer_pattern") != "no_live_contact_or_voicemail"
        and clean(row.get("ideal_answer_manager_sanitized"))
    ]
    candidates = sorted(candidates, key=lambda row: (-clamp_int(row.get("overall_quality_score"), 0, 100, 0), clean(row.get("llm_customer_signal_type")), clean(row.get("moment_id"))))
    return [example_row(row) for row in candidates[:limit]]


def build_rop_coaching_queue(reviews: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    candidates = [
        row
        for row in reviews
        if is_trusted_llm_review(row)
        and (
            clamp_int(row.get("overall_quality_score"), 0, 100, 0) < 60
            or row.get("answer_pattern") in {"vague_or_missing_next_step", "service_answer_without_closure", "price_payment_answer_too_weak"}
        )
    ]
    candidates = sorted(candidates, key=lambda row: (clamp_int(row.get("overall_quality_score"), 0, 100, 0), clean(row.get("manager_name")), clean(row.get("moment_id"))))
    return [coaching_row(row) for row in candidates[:limit]]


def build_bot_knowledge_seeds(reviews: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    candidates = [
        row
        for row in reviews
        if is_trusted_llm_review(row)
        and row.get("bot_seed_status") in {"ready_for_bot_draft", "needs_rop_validation"}
        and row.get("answer_pattern") != "no_live_contact_or_voicemail"
        and clean(row.get("bot_safe_answer"))
        and row.get("bot_safety_status") not in BOT_BLOCKING_SAFETY_STATUSES
    ]
    candidates = sorted(candidates, key=lambda row: (row.get("bot_seed_status") != "ready_for_bot_draft", -clamp_int(row.get("overall_quality_score"), 0, 100, 0), clean(row.get("llm_customer_signal_type"))))
    seen: set[tuple[str, str]] = set()
    rows: list[dict[str, Any]] = []
    for row in candidates:
        key = (clean(row.get("llm_customer_signal_type")), normalize_topic(clean(row.get("customer_question"))))
        if key in seen:
            continue
        seen.add(key)
        rows.append(
            {
                "Статус для бота": row.get("bot_seed_status_ru", ""),
                "Код статуса для бота": row.get("bot_seed_status", ""),
                "Сигнал клиента": row.get("signal_ru", ""),
                "Код сигнала": row.get("llm_customer_signal_type", ""),
                "Стадия": row.get("stage_ru", ""),
                "Паттерн ответа": row.get("answer_pattern_ru", ""),
                "Код паттерна": row.get("answer_pattern", ""),
                "Пример вопроса клиента": row.get("customer_question_sanitized", row.get("customer_question", "")),
                "Черновик идеального ответа": row.get("bot_safe_answer", ""),
                "Безопасный ответ для бота": row.get("bot_safe_answer", ""),
                "Идеальный ответ для менеджера": row.get("ideal_answer_manager_sanitized", ""),
                "Статус sanitizer": row.get("bot_safety_status_ru", ""),
                "Флаги sanitizer": row.get("sanitizer_flags", ""),
                "Риск бренда": row.get("brand_risk_flag", ""),
                "Риск цены/скидки": row.get("money_or_discount_flag", ""),
                "Риск рассрочки": row.get("installment_flag", ""),
                "Риск договора/возврата": row.get("legal_or_refund_flag", ""),
                "Риск срока/обещания": row.get("deadline_or_promise_flag", ""),
                "Риск персональных данных": row.get("personal_data_flag", ""),
                "Когда не использовать": row.get("avoid_using_when", ""),
                "Оценка": row.get("overall_quality_score", ""),
                "Итог сделки": row.get("final_outcome_ru", ""),
                "Код итога": row.get("final_outcome_label", ""),
                "Цитата клиента": row.get("customer_quote_sanitized", row.get("customer_quote", "")),
                "Цитата менеджера": row.get("manager_quote_sanitized", row.get("manager_quote", "")),
                "Ограничение данных": row.get("data_scope_note", ""),
                "ID момента": row.get("moment_id", ""),
            }
        )
        if len(rows) >= limit:
            break
    return rows


def build_llm_refresh_queue(reviews: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in reviews:
        if is_trusted_llm_review(row):
            continue
        rows.append(
            {
                "Статус проверки": row.get("review_trust_status_ru", ""),
                "Провайдер текущей строки": row.get("provider", ""),
                "Источник ревью": row.get("review_source", ""),
                "ID момента": row.get("moment_id", ""),
                "Телефон": row.get("phone", ""),
                "Дата звонка": row.get("started_at", ""),
                "Менеджер": row.get("manager_name", ""),
                "Тип звонка": row.get("call_type", ""),
                "Use case": row.get("extraction_use_case", ""),
                "Сигнал клиента": row.get("signal_ru", ""),
                "Стадия": row.get("stage_ru", ""),
                "Итог сделки": row.get("final_outcome_ru", ""),
                "Вопрос клиента": row.get("customer_question", ""),
                "Ответ менеджера": row.get("manager_answer", ""),
                "Причина": "Нужно повторить live LLM-review перед использованием в KB/боте/коучинге.",
                "Файл звонка": row.get("source_filename", ""),
            }
        )
    return rows


def build_manager_summary(reviews: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for manager, group_rows in sorted(group_by(reviews, "manager_name").items()):
        if not manager:
            continue
        scores = [clamp_int(row.get("overall_quality_score"), 0, 100, 0) for row in group_rows]
        rows.append(
            {
                "Менеджер": manager,
                "Количество": len(group_rows),
                "Среднее качество": avg(scores),
                "Высокое качество": sum(1 for score in scores if score >= 75),
                "Низкое качество": sum(1 for score in scores if score < 55),
                "Топ сигналы": top_labeled_values((row.get("llm_customer_signal_type") for row in group_rows), signal_label_ru, 4),
                "Топ паттерны": top_labeled_values((row.get("answer_pattern") for row in group_rows), answer_pattern_label_ru, 4),
                "Главный фокус разбора": manager_focus(group_rows),
                "Ограничение данных": "Оценка только по звонкам; переписка не учтена.",
            }
        )
    return sorted(rows, key=lambda row: (-int(row["Количество"]), float(row["Среднее качество"])))


def build_outcome_lens(reviews: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for outcome, group_rows in sorted(group_by(reviews, "outcome_group").items()):
        scores = [clamp_int(row.get("overall_quality_score"), 0, 100, 0) for row in group_rows]
        rows.append(
            {
                "Группа итога": outcome_group_label_ru(outcome),
                "Код группы итога": outcome,
                "Количество": len(group_rows),
                "Среднее качество": avg(scores),
                "Высокое качество": sum(1 for score in scores if score >= 75),
                "Низкое качество": sum(1 for score in scores if score < 55),
                "Топ сигналы": top_labeled_values((row.get("llm_customer_signal_type") for row in group_rows), signal_label_ru, 5),
                "Топ паттерны": top_labeled_values((row.get("answer_pattern") for row in group_rows), answer_pattern_label_ru, 5),
                "Интерпретация": outcome_interpretation(outcome),
            }
        )
    return sorted(rows, key=lambda row: (-int(row["Количество"]), str(row["Группа итога"])))


def build_rop_brief(
    summary: dict[str, Any],
    signal_summary: list[dict[str, Any]],
    pattern_matrix: list[dict[str, Any]],
    rop_coaching: list[dict[str, Any]],
    bot_seeds: list[dict[str, Any]],
    manager_summary: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    counts = summary.get("counts", {})
    usefulness = counts.get("by_commercial_usefulness", {})
    bot_status = counts.get("by_bot_seed_status", {})
    pattern_counts = counts.get("by_answer_pattern", {})
    rows = [
        {
            "Раздел": "Как читать",
            "Показатель": "Ограничение данных",
            "Значение": "",
            "Комментарий": "Выводы сделаны только по звонкам. Если менеджер продолжил работу в WhatsApp/Telegram/email, здесь это не видно; такие строки нужно считать гипотезой для проверки, а не обвинением менеджера.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Проверено звонковых моментов",
            "Значение": summary.get("totals", {}).get("reviews", 0),
            "Комментарий": "Объем текущей версии базы знаний продаж.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Доверенных live/GPT-ревью",
            "Значение": summary.get("llm_review", {}).get("trusted_llm_reviews", 0),
            "Комментарий": "Только эти строки используются для лучших ответов, бота, паттернов и очереди РОПа.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Ждут live LLM-refresh",
            "Значение": summary.get("llm_review", {}).get("needs_live_llm_refresh", 0),
            "Комментарий": "Эти строки не попадают в базу бота и лучшие ответы до повторного live-ревью.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Среднее качество ответа",
            "Значение": summary.get("quality", {}).get("avg_quality_score", 0),
            "Комментарий": "Средняя оценка ИИ по рубрике качества продаж/сервиса.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Кандидаты в хорошие примеры для скриптов",
            "Значение": usefulness.get("playbook_candidate", 0),
            "Комментарий": "Кандидаты в хорошие ответы после ручной проверки.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Риск потери выручки",
            "Значение": usefulness.get("revenue_leakage_risk", 0),
            "Комментарий": "Слабые ответы в цепочках с потерей/реактивацией.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Черновики для базы ответов бота",
            "Значение": bot_status.get("ready_for_bot_draft", 0),
            "Комментарий": "Черновики ответов для будущей базы Telegram-бота.",
        },
        {
            "Раздел": "Ключевые показатели",
            "Показатель": "Безопасных bot-safe ответов",
            "Значение": summary.get("sanitizer", {}).get("bot_safe_answer_rows", 0),
            "Комментарий": "Ответы, где sanitizer убрал брендовые/финансовые/сроковые/персональные риски.",
        },
    ]
    for pattern, count in sorted(pattern_counts.items(), key=lambda item: (-int(item[1]), str(item[0])))[:7]:
        rows.append(
            {
                "Раздел": "Топ паттерны",
                "Показатель": answer_pattern_label_ru(pattern),
                "Значение": count,
                "Комментарий": pattern_level_recommendation(pattern, "", 0),
            }
        )
    for row in signal_summary[:7]:
        rows.append(
            {
                "Раздел": "Сигналы клиентов",
                "Показатель": row.get("Сигнал клиента", ""),
                "Значение": row.get("Количество", ""),
                "Комментарий": row.get("Рекомендация РОПу", ""),
            }
        )
    for row in rop_coaching[:10]:
        rows.append(
            {
                "Раздел": "Очередь РОПа",
                "Показатель": f"{row.get('Приоритет', '')}: {row.get('Менеджер', '')}",
                "Значение": row.get("Оценка", ""),
                "Комментарий": row.get("Что сделать РОПу", ""),
            }
        )
    for row in manager_summary[:8]:
        rows.append(
            {
                "Раздел": "Менеджеры",
                "Показатель": row.get("Менеджер", ""),
                "Значение": row.get("Среднее качество", ""),
                "Комментарий": row.get("Главный фокус разбора", ""),
            }
        )
    for row in bot_seeds[:10]:
        rows.append(
            {
                "Раздел": "Черновики для бота",
                "Показатель": row.get("Сигнал клиента", ""),
                "Значение": row.get("Оценка", ""),
                "Комментарий": row.get("Безопасный ответ для бота", row.get("Черновик идеального ответа", "")),
            }
        )
    if pattern_matrix:
        rows.append(
            {
                "Раздел": "Следующий аналитический шаг",
                "Показатель": "Расширить выборку с 160 до 2734 моментов",
                "Значение": "",
                "Комментарий": "После ручной проверки РОПом можно прогнать всю пилотную выборку и построить более устойчивые связи между вопросами, ответами и итогами сделок.",
            }
        )
    return rows


def example_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "Сигнал клиента": row.get("signal_ru", ""),
        "Код сигнала": row.get("llm_customer_signal_type", ""),
        "Стадия": row.get("stage_ru", ""),
        "Паттерн ответа": row.get("answer_pattern_ru", ""),
        "Код паттерна": row.get("answer_pattern", ""),
        "Итог сделки": row.get("final_outcome_ru", ""),
        "Код итога": row.get("final_outcome_label", ""),
        "Оценка": row.get("overall_quality_score", ""),
        "Менеджер": row.get("manager_name", ""),
        "Вопрос клиента": row.get("customer_question_sanitized", row.get("customer_question", "")),
        "Ответ менеджера": row.get("manager_answer", ""),
        "Идеальный ответ": row.get("ideal_answer_manager_sanitized", row.get("ideal_answer_example", "")),
        "Идеальный ответ для менеджера": row.get("ideal_answer_manager_sanitized", ""),
        "Безопасный ответ для бота": row.get("bot_safe_answer", ""),
        "Статус sanitizer": row.get("bot_safety_status_ru", ""),
        "Флаги sanitizer": row.get("sanitizer_flags", ""),
        "Риск бренда": row.get("brand_risk_flag", ""),
        "Риск цены/скидки": row.get("money_or_discount_flag", ""),
        "Риск рассрочки": row.get("installment_flag", ""),
        "Риск договора/возврата": row.get("legal_or_refund_flag", ""),
        "Риск срока/обещания": row.get("deadline_or_promise_flag", ""),
        "Риск персональных данных": row.get("personal_data_flag", ""),
        "Что хорошо": row.get("what_manager_did_well", ""),
        "Риски": row.get("risk_flags", ""),
        "Ограничение данных": row.get("data_scope_note", ""),
        "ID момента": row.get("moment_id", ""),
        "Файл звонка": row.get("source_filename", ""),
    }


def coaching_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "Приоритет": coaching_priority(row),
        "Менеджер": row.get("manager_name", ""),
        "Сигнал клиента": row.get("signal_ru", ""),
        "Код сигнала": row.get("llm_customer_signal_type", ""),
        "Стадия": row.get("stage_ru", ""),
        "Паттерн ответа": row.get("answer_pattern_ru", ""),
        "Код паттерна": row.get("answer_pattern", ""),
        "Итог сделки": row.get("final_outcome_ru", ""),
        "Код итога": row.get("final_outcome_label", ""),
        "Оценка": row.get("overall_quality_score", ""),
        "Что упущено": row.get("what_manager_missed", ""),
        "Риски": row.get("risk_flags", ""),
        "Что сделать РОПу": row.get("rop_action", ""),
        "Ограничение данных": row.get("data_scope_note", ""),
        "Вопрос клиента": row.get("customer_question_sanitized", row.get("customer_question", "")),
        "Ответ менеджера": row.get("manager_answer", ""),
        "Идеальный ответ": row.get("ideal_answer_manager_sanitized", row.get("ideal_answer_example", "")),
        "Идеальный ответ для менеджера": row.get("ideal_answer_manager_sanitized", ""),
        "Безопасный ответ для бота": row.get("bot_safe_answer", ""),
        "Статус sanitizer": row.get("bot_safety_status_ru", ""),
        "Флаги sanitizer": row.get("sanitizer_flags", ""),
        "Риск бренда": row.get("brand_risk_flag", ""),
        "Риск цены/скидки": row.get("money_or_discount_flag", ""),
        "Риск рассрочки": row.get("installment_flag", ""),
        "Риск договора/возврата": row.get("legal_or_refund_flag", ""),
        "Риск срока/обещания": row.get("deadline_or_promise_flag", ""),
        "Риск персональных данных": row.get("personal_data_flag", ""),
        "ID момента": row.get("moment_id", ""),
        "Файл звонка": row.get("source_filename", ""),
    }


def write_outputs(
    out_root: Path,
    summary: dict[str, Any],
    rop_brief: list[dict[str, Any]],
    signal_summary: list[dict[str, Any]],
    pattern_matrix: list[dict[str, Any]],
    best_answers: list[dict[str, Any]],
    rop_coaching: list[dict[str, Any]],
    bot_seeds: list[dict[str, Any]],
    manager_summary: list[dict[str, Any]],
    outcome_lens: list[dict[str, Any]],
    llm_refresh_queue: list[dict[str, Any]],
    enriched_reviews: list[dict[str, Any]],
) -> dict[str, Path]:
    outputs = {
        "signal_summary_csv": out_root / "signal_summary.csv",
        "pattern_matrix_csv": out_root / "pattern_matrix.csv",
        "best_answers_csv": out_root / "best_answers.csv",
        "rop_coaching_csv": out_root / "rop_coaching_queue.csv",
        "bot_seeds_csv": out_root / "bot_knowledge_seeds.csv",
        "manager_summary_csv": out_root / "manager_summary.csv",
        "outcome_lens_csv": out_root / "outcome_lens.csv",
        "llm_refresh_queue_csv": out_root / "llm_refresh_queue.csv",
        "enriched_reviews_csv": out_root / "enriched_reviews.csv",
    }
    write_csv(outputs["signal_summary_csv"], signal_summary)
    write_csv(outputs["pattern_matrix_csv"], pattern_matrix)
    write_csv(outputs["best_answers_csv"], best_answers)
    write_csv(outputs["rop_coaching_csv"], rop_coaching)
    write_csv(outputs["bot_seeds_csv"], bot_seeds)
    write_csv(outputs["manager_summary_csv"], manager_summary)
    write_csv(outputs["outcome_lens_csv"], outcome_lens)
    write_csv(outputs["llm_refresh_queue_csv"], llm_refresh_queue)
    write_csv(outputs["enriched_reviews_csv"], enriched_reviews)
    xlsx_path = out_root / "sales_insight_knowledge_base.xlsx"
    write_xlsx(
        xlsx_path,
        summary,
        rop_brief,
        signal_summary,
        pattern_matrix,
        best_answers,
        rop_coaching,
        bot_seeds,
        manager_summary,
        outcome_lens,
        llm_refresh_queue,
        enriched_reviews,
    )
    outputs["xlsx"] = xlsx_path
    return outputs


def write_xlsx(
    path: Path,
    summary: dict[str, Any],
    rop_brief: list[dict[str, Any]],
    signal_summary: list[dict[str, Any]],
    pattern_matrix: list[dict[str, Any]],
    best_answers: list[dict[str, Any]],
    rop_coaching: list[dict[str, Any]],
    bot_seeds: list[dict[str, Any]],
    manager_summary: list[dict[str, Any]],
    outcome_lens: list[dict[str, Any]],
    llm_refresh_queue: list[dict[str, Any]],
    enriched_reviews: list[dict[str, Any]],
) -> None:
    import pandas as pd

    summary_rows: list[dict[str, Any]] = []
    for section in ("totals", "quality"):
        for key, value in summary.get(section, {}).items():
            summary_rows.append({"Раздел": summary_section_label_ru(section), "Метрика": summary_metric_label_ru(key), "Значение": value})
    for key, value in summary.get("llm_review", {}).items():
        if isinstance(value, dict):
            continue
        summary_rows.append({"Раздел": summary_section_label_ru("llm_review"), "Метрика": summary_metric_label_ru(key), "Значение": value})
    for key, value in summary.get("sanitizer", {}).items():
        if isinstance(value, dict):
            continue
        summary_rows.append({"Раздел": summary_section_label_ru("sanitizer"), "Метрика": summary_metric_label_ru(key), "Значение": value})
    for group, counts in summary.get("counts", {}).items():
        for label, count in counts.items():
            summary_rows.append({"Раздел": summary_section_label_ru(group), "Метрика": summary_count_label_ru(group, label), "Значение": count})
    for note in summary.get("audit_notes", []):
        summary_rows.append({"Раздел": "Ограничения и аудит", "Метрика": "Комментарий", "Значение": note})

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(rop_brief).to_excel(writer, sheet_name="Сводка РОПа", index=False)
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Итоги", index=False)
        pd.DataFrame(signal_summary).to_excel(writer, sheet_name="Сигналы", index=False)
        pd.DataFrame(pattern_matrix).to_excel(writer, sheet_name="Паттерны", index=False)
        pd.DataFrame(best_answers).to_excel(writer, sheet_name="Лучшие ответы", index=False)
        pd.DataFrame(rop_coaching).to_excel(writer, sheet_name="Разбор РОПа", index=False)
        pd.DataFrame(bot_seeds).to_excel(writer, sheet_name="Черновики бота", index=False)
        pd.DataFrame(manager_summary).to_excel(writer, sheet_name="Менеджеры", index=False)
        pd.DataFrame(outcome_lens).to_excel(writer, sheet_name="Итоги сделок", index=False)
        pd.DataFrame(llm_refresh_queue).to_excel(writer, sheet_name="LLM refresh", index=False)
        pd.DataFrame(enriched_reviews).to_excel(writer, sheet_name="Исходные строки", index=False)
        style_workbook(writer.book)


def style_workbook(workbook: Any) -> None:
    from openpyxl.styles import Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for column_cells in sheet.columns:
            max_len = 0
            column = column_cells[0].column_letter
            for cell in column_cells[:200]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
                if len(value) > 70:
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = True
                    alignment.vertical = "top"
                    cell.alignment = alignment
            sheet.column_dimensions[column].width = min(max(max_len + 2, 10), 62)


def signal_level_recommendation(signal: str, rows: list[dict[str, Any]]) -> str:
    low_share = sum(1 for row in rows if clamp_int(row.get("overall_quality_score"), 0, 100, 0) < 55) / max(len(rows), 1)
    if signal in {"price_question", "price_objection", "discount_or_installment_question"}:
        return "Собрать единый блок отработки цены: ценность, формат, рассрочка/скидка, точный следующий шаг."
    if signal in {"schedule_question", "callback_request", "materials_request"}:
        return "Проверить дисциплину повторного контакта: дата, время и канал должны фиксироваться в CRM."
    if signal in {"complaint_or_service_risk", "technical_or_access_issue", "existing_client_progress"}:
        return "Сделать сценарий восстановления сервиса: признание проблемы, срок решения, контроль закрытия."
    if low_share >= 0.35:
        return "Высокая доля слабых ответов: разобрать 3-5 звонков с РОПом и обновить скрипт."
    return "Использовать лучшие примеры как материал для скриптов и обучения менеджеров."


def pattern_level_recommendation(pattern: str, outcome: str, avg_score: float) -> str:
    if pattern == "vague_or_missing_next_step":
        return "Ввести обязательный конкретный следующий шаг; это процессная потеря, а не проблема знаний."
    if pattern == "no_live_contact_or_voicemail":
        return "Настроить правило повторного касания и альтернативный канал после недозвона."
    if pattern == "service_answer_without_closure":
        return "Закрывать сервисный запрос сроком, ответственным и проверкой результата."
    if avg_score >= 75 and outcome in {"paid_or_payment_path", "follow_up_opportunity"}:
        return "Кандидат в базу лучших ответов: проверить вручную и использовать в обучении/боте."
    return "Использовать как сегмент для ручной проверки и уточнения скрипта."


def manager_focus(rows: list[dict[str, Any]]) -> str:
    low_rows = [row for row in rows if clamp_int(row.get("overall_quality_score"), 0, 100, 0) < 60]
    source = low_rows or rows
    patterns = Counter(clean(row.get("answer_pattern")) for row in source if clean(row.get("answer_pattern")))
    if not patterns:
        return "Недостаточно данных."
    pattern, _count = patterns.most_common(1)[0]
    return pattern_level_recommendation(pattern, "", 0)


def coaching_priority(row: dict[str, Any]) -> str:
    score = clamp_int(row.get("overall_quality_score"), 0, 100, 0)
    outcome = clean(row.get("outcome_group"))
    if score < 45 and outcome in {"lost_or_churn", "follow_up_opportunity"}:
        return "P0 риск потери выручки"
    if score < 55:
        return "P1 слабый ответ"
    if row.get("answer_pattern") == "vague_or_missing_next_step":
        return "P1 процесс: нет точного следующего шага"
    return "P2 обучение"


def outcome_interpretation(outcome: str) -> str:
    return {
        "paid_or_payment_path": "Искать, какие ответы чаще сопровождают оплату или намерение оплатить.",
        "follow_up_opportunity": "Это зона потенциальных денег: важны повторный контакт и реактивация.",
        "retention_or_service": "Это не всегда новая продажа, но влияет на продление и доверие.",
        "lost_or_churn": "Искать повторяющиеся потери: цена, формат, следующий шаг, сервисные сбои.",
        "manual_or_mixed": "Нужна ручная проверка: смешанный или спорный итог.",
    }.get(outcome, "Недостаточно уверенный итог, использовать осторожно.")


def normalize_topic(text: str) -> str:
    text = text.lower()
    for pattern, label in (
        (r"цен|стоимост|скидк|рассроч|маткапитал", "price"),
        (r"распис|врем|день|очно|онлайн", "schedule_format"),
        (r"программ|курс|групп|уровн|егэ|огэ|олимпиад", "program_level"),
        (r"оплат|счет|счёт|квитанц|договор", "payment_docs"),
        (r"личн|кабинет|ссылк|доступ|чат", "tech_access"),
    ):
        if re.search(pattern, text):
            return label
    return clean(text)[:60]


def best_row(rows: list[dict[str, Any]]) -> dict[str, Any]:
    if not rows:
        return {}
    return sorted(rows, key=lambda row: (-clamp_int(row.get("overall_quality_score"), 0, 100, 0), -clamp_float(row.get("extraction_confidence"), 0, 1, 0), clean(row.get("moment_id"))))[0]


def top_values(values: Iterable[Any], limit: int) -> str:
    counter = Counter(clean(value) for value in values if clean(value))
    return " | ".join(f"{key}: {count}" for key, count in counter.most_common(limit))


def top_labeled_values(values: Iterable[Any], label_func: Any, limit: int) -> str:
    counter = Counter(clean(value) for value in values if clean(value))
    return " | ".join(f"{label_func(key)}: {count}" for key, count in counter.most_common(limit))


def sanitizer_flag_counter(rows: list[dict[str, Any]]) -> Counter[str]:
    counter: Counter[str] = Counter()
    for row in rows:
        for flag in clean(row.get("sanitizer_flags")).split("|"):
            flag = flag.strip()
            if flag:
                counter[flag] += 1
    return counter


def summary_count_label_ru(group: str, label: str) -> str:
    if group == "by_signal":
        return signal_label_ru(label)
    if group == "by_answer_pattern":
        return answer_pattern_label_ru(label)
    if group == "by_outcome_group":
        return outcome_group_label_ru(label)
    if group == "by_bot_seed_status":
        return bot_status_label_ru(label)
    if group == "by_commercial_usefulness":
        return usefulness_label_ru(label)
    if group == "by_bot_safety_status":
        return bot_safety_label_ru(label)
    return clean(label)


def summary_section_label_ru(value: Any) -> str:
    code = clean(value)
    return {
        "totals": "Общие итоги",
        "quality": "Качество ответов",
        "llm_review": "Доверенность LLM-ревью",
        "sanitizer": "Sanitizer для бота/РОПа",
        "by_signal": "Сигналы клиентов",
        "by_answer_pattern": "Паттерны ответов",
        "by_outcome_group": "Итоги сделок",
        "by_bot_seed_status": "Готовность для бота",
        "by_commercial_usefulness": "Коммерческая полезность",
    }.get(code, code)


def summary_metric_label_ru(value: Any) -> str:
    code = clean(value)
    return {
        "reviews": "Проверено моментов",
        "unique_phones": "Уникальных телефонов",
        "unique_managers": "Уникальных менеджеров",
        "signals": "Типов сигналов клиента",
        "answer_patterns": "Типов паттернов ответа",
        "avg_quality_score": "Средняя оценка качества",
        "high_quality_count": "Ответов высокого качества",
        "low_quality_count": "Ответов низкого качества",
        "trusted_llm_reviews": "Доверенных live/GPT-ревью",
        "needs_live_llm_refresh": "Нужно повторить live LLM-refresh",
        "trusted_share": "Доля доверенных LLM-ревью",
        "bot_ready_or_validation_rows": "Bot-ready / needs validation строк",
        "bot_safe_answer_rows": "Строк с безопасным ответом для бота",
        "bot_safety_blocked": "Заблокировано sanitizer-ом",
    }.get(code, code)


def signal_label_ru(value: Any) -> str:
    code = clean(value) or "unknown"
    return SIGNAL_LABELS_RU.get(code, code)


def stage_label_ru(value: Any) -> str:
    code = clean(value) or "unknown"
    return STAGE_LABELS_RU.get(code, code)


def answer_pattern_label_ru(value: Any) -> str:
    code = clean(value) or "generic_answer"
    return ANSWER_PATTERN_LABELS_RU.get(code, code)


def outcome_group_label_ru(value: Any) -> str:
    code = clean(value) or "unknown_or_other"
    return OUTCOME_GROUP_LABELS_RU.get(code, code)


def final_outcome_label_ru(value: Any) -> str:
    code = clean(value) or "unknown"
    return FINAL_OUTCOME_LABELS_RU.get(code, code)


def bot_status_label_ru(value: Any) -> str:
    code = clean(value) or "not_ready"
    return BOT_STATUS_LABELS_RU.get(code, code)


def usefulness_label_ru(value: Any) -> str:
    code = clean(value) or "useful_context"
    return USEFULNESS_LABELS_RU.get(code, code)


def review_trust_label_ru(value: Any) -> str:
    code = clean(value) or "needs_live_llm_refresh"
    return REVIEW_TRUST_LABELS_RU.get(code, code)


def bot_safety_label_ru(value: Any) -> str:
    code = clean(value) or "empty"
    return BOT_SAFETY_LABELS_RU.get(code, code)


def share(rows: list[dict[str, Any]], key: str, value: str) -> float:
    if not rows:
        return 0.0
    return round(sum(1 for row in rows if clean(row.get(key)) == value) / len(rows), 3)


def avg(values: list[int]) -> float:
    return round(sum(values) / len(values), 2) if values else 0.0


def group_by(rows: list[dict[str, Any]], key: str) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[clean(row.get(key))].append(row)
    return grouped


def read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return re.sub(r"\s+", " ", text)


def clamp_int(value: Any, low: int, high: int, default: int) -> int:
    try:
        parsed = int(round(float(str(value).strip())))
    except (TypeError, ValueError):
        return default
    return max(low, min(high, parsed))


def clamp_float(value: Any, low: float, high: float, default: float) -> float:
    try:
        parsed = float(str(value).strip())
    except (TypeError, ValueError):
        return default
    return max(low, min(high, parsed))


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build sales insight knowledge-base aggregates from LLM review rows.")
    parser.add_argument("--project-root", default=".")
    parser.add_argument("--reviews-csv", default="stable_runtime/pilot_sales_moment_llm_review_20260507_codex_batch160/reviews.csv")
    parser.add_argument("--out-root", default="stable_runtime/sales_insight_knowledge_base_20260507")
    parser.add_argument("--min-group-count", type=int, default=2)
    parser.add_argument("--top-examples", type=int, default=120)
    return parser.parse_args(argv)


def config_from_args(args: argparse.Namespace) -> KnowledgeBaseConfig:
    project_root = Path(args.project_root).expanduser().resolve()
    return KnowledgeBaseConfig(
        project_root=project_root,
        reviews_csv=(project_root / args.reviews_csv).resolve(),
        out_root=(project_root / args.out_root).resolve(),
        min_group_count=max(1, int(args.min_group_count)),
        top_examples=max(1, int(args.top_examples)),
    )


__all__ = [
    "KnowledgeBaseConfig",
    "build_sales_insight_knowledge_base",
    "classify_answer_pattern",
    "commercial_usefulness",
    "config_from_args",
    "enrich_review_row",
    "outcome_group",
    "parse_args",
    "quality_band",
    "sanitizer_flag_counter",
]
